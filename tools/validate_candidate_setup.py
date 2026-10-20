#!/usr/bin/env python3
"""
Candidate setup validator.

Run this BEFORE spending any money on model calls:

    python3 tools/validate_candidate_setup.py

Checks that every required private candidate file exists, parses, and has
the shapes the deterministic pipeline actually expects (see
docs/candidate-data-guide.md). Makes NO network/API/model calls.

Exit code 0 means every hard check passed (warnings may still be printed).
Exit code 1 means at least one hard check failed.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MASTER = ROOT / "master"
REFERENCE = ROOT / "reference"

REQUIRED_MASTER_FILES = {
    "candidate_facts.md": MASTER / "candidate_facts.md",
    "candidate_profile.json": MASTER / "candidate_profile.json",
    "evidence_bank.json": MASTER / "evidence_bank.json",
    "metrics_registry.json": MASTER / "metrics_registry.json",
    "eligibility_facts.json": MASTER / "eligibility_facts.json",
    "rejections.json": MASTER / "rejections.json",
    "canonical_resume_entries.md": MASTER / "canonical_resume_entries.md",
    "runtime_resume_context.json": MASTER / "runtime" / "runtime_resume_context.json",
}

REQUIRED_TEMPLATES = {
    "resume template": REFERENCE / "resume_template.docx",
    "cover-letter template": REFERENCE / "cover_letter_template.docx",
}

REQUIRED_LAYOUTS = {
    "resume template layout": REFERENCE / "resume_template_layout.json",
    "cover-letter template layout": REFERENCE / "cover_letter_template_layout.json",
}

EXAMPLE_PLACEHOLDER_VALUES = {
    "Candidate Example",
    "candidate@example.com",
    "(555) 010-2000",
    "linkedin.com/in/candidate-example",
    "github.com/example-user",
}


class Report:
    def __init__(self):
        self.failures: list[str] = []
        self.warnings: list[str] = []

    def fail(self, message: str) -> None:
        self.failures.append(message)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def ok(self) -> bool:
        return not self.failures


def load_json(path: Path, report: Report) -> dict | None:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        report.fail(f"Missing file: {path.relative_to(ROOT)}")
    except json.JSONDecodeError as exc:
        report.fail(f"Invalid JSON in {path.relative_to(ROOT)}: {exc}")
    return None


def check_files_exist(report: Report) -> None:
    for label, path in REQUIRED_MASTER_FILES.items():
        if not path.exists():
            report.fail(f"Missing required candidate file: {path.relative_to(ROOT)}")


def check_runtime_context(report: Report) -> dict | None:
    path = REQUIRED_MASTER_FILES["runtime_resume_context.json"]
    if not path.exists():
        return None

    data = load_json(path, report)
    if data is None:
        return None

    for key in ("schema_version", "candidate", "education", "experience", "projects", "skills", "truth_constraints"):
        if key not in data:
            report.fail(f"runtime_resume_context.json missing top-level key: {key}")

    candidate = data.get("candidate", {}) or {}
    for field in ("full_name", "email"):
        value = candidate.get(field)
        if not value or not str(value).strip():
            report.fail(f"runtime_resume_context.json candidate.{field} is missing/empty")
        elif value in EXAMPLE_PLACEHOLDER_VALUES:
            report.warn(f"runtime_resume_context.json candidate.{field} still has the example value {value!r}")

    education = data.get("education", []) or []
    if not education:
        report.fail("runtime_resume_context.json must have at least one education entry")

    experience = data.get("experience", []) or []
    if not experience:
        report.fail("runtime_resume_context.json must have at least one experience entry")

    seen_experience_ids = set()
    tailorable_experience_ids = set()
    for i, entry in enumerate(experience):
        entry_id = entry.get("id")
        if not entry_id:
            report.fail(f"experience[{i}] is missing required field: id")
        elif entry_id in seen_experience_ids:
            report.fail(f"experience[{i}] has a duplicate id: {entry_id}")
        else:
            seen_experience_ids.add(entry_id)

        mode = entry.get("tailoring_mode")
        if mode not in ("protected", "tailorable"):
            report.fail(
                f"experience[{i}] (id={entry_id}) tailoring_mode must be "
                f"'protected' or 'tailorable'; got {mode!r}"
            )
        elif mode == "protected" and not entry.get("canonical_bullets"):
            report.fail(f"experience[{i}] (id={entry_id}) is protected but has no canonical_bullets")
        elif mode == "tailorable":
            tailorable_experience_ids.add(entry_id)
            if not entry.get("alternate_bullets_for_dynamic_tailoring"):
                report.fail(
                    f"experience[{i}] (id={entry_id}) is tailorable but has no "
                    "alternate_bullets_for_dynamic_tailoring"
                )

        if not entry.get("employer"):
            report.fail(f"experience[{i}] (id={entry_id}) is missing employer")

    if not tailorable_experience_ids:
        report.warn(
            "No experience entry has tailoring_mode 'tailorable'. The pipeline can still "
            "run, but no experience bullet will ever be dynamically tailored per job."
        )

    projects = data.get("projects", []) or []
    if not projects:
        report.fail("runtime_resume_context.json must have at least one project entry")

    seen_project_ids = set()
    selectable_count = 0
    for i, entry in enumerate(projects):
        entry_id = entry.get("id")
        if not entry_id:
            report.fail(f"projects[{i}] is missing required field: id")
        elif entry_id in seen_project_ids:
            report.fail(f"projects[{i}] has a duplicate id: {entry_id}")
        else:
            seen_project_ids.add(entry_id)

        status = entry.get("project_status")
        if status not in ("required", "selectable"):
            report.fail(
                f"projects[{i}] (id={entry_id}) project_status must be "
                f"'required' or 'selectable'; got {status!r}"
            )

        mode = entry.get("tailoring_mode")
        if mode not in ("protected", "tailorable"):
            report.fail(
                f"projects[{i}] (id={entry_id}) tailoring_mode must be "
                f"'protected' or 'tailorable'; got {mode!r}"
            )
        elif mode == "protected" and not entry.get("canonical_bullets"):
            report.fail(f"projects[{i}] (id={entry_id}) is protected but has no canonical_bullets")
        elif mode == "tailorable":
            if not entry.get("alternate_bullets_for_dynamic_tailoring"):
                report.fail(
                    f"projects[{i}] (id={entry_id}) is tailorable but has no "
                    "alternate_bullets_for_dynamic_tailoring"
                )

        if status == "selectable":
            selectable_count += 1

        if not entry.get("name"):
            report.fail(f"projects[{i}] (id={entry_id}) is missing name")

    if selectable_count == 0:
        report.fail(
            "At least one project must have project_status == 'selectable' "
            "(tools/resume/v8/select_project.py has no candidate pool to choose from otherwise)."
        )

    skills = data.get("skills", {}) or {}
    if not skills.get("static_resume_lines"):
        report.warn("skills.static_resume_lines is empty -- the Skills section will render empty")

    if not data.get("truth_constraints"):
        report.warn("truth_constraints is empty -- consider adding at least one")

    return data


def check_project_selection_pool(report: Report, runtime_context: dict) -> None:
    evidence_path = MASTER / "evidence_bank.json"
    evidence = load_json(evidence_path, report)
    if evidence is None:
        return

    entities = {str(item.get("entity", "")) for item in evidence.get("items", []) if isinstance(item, dict)}
    project_names = {p.get("name") for p in runtime_context.get("projects", []) if p.get("project_status") == "selectable"}

    for name in project_names:
        if name and name not in entities:
            report.warn(
                f"Selectable project {name!r} has no evidence_bank.json items with "
                f"entity == {name!r} -- it will never be able to demonstrate incremental "
                "JD coverage during selection."
            )


def check_metrics_registry(report: Report, runtime_context: dict) -> None:
    path = MASTER / "metrics_registry.json"
    data = load_json(path, report)
    if data is None:
        return

    valid_ids = {e.get("id") for e in runtime_context.get("experience", [])}
    valid_ids |= {p.get("id") for p in runtime_context.get("projects", [])}

    for i, metric in enumerate(data.get("metrics", []) or []):
        owner_id = metric.get("owner_id")
        if owner_id and owner_id not in valid_ids:
            report.warn(
                f"metrics_registry.json metrics[{i}] owner_id {owner_id!r} does not match "
                "any experience/project id in runtime_resume_context.json"
            )
        if not str(metric.get("statement", "")).strip():
            report.fail(f"metrics_registry.json metrics[{i}] is missing a statement")


def check_layouts(report: Report, runtime_context: dict) -> None:
    for label, path in REQUIRED_TEMPLATES.items():
        if not path.exists():
            report.fail(f"Missing {label}: {path.relative_to(ROOT)}")

    for label, path in REQUIRED_LAYOUTS.items():
        if not path.exists():
            report.fail(
                f"Missing {label}: {path.relative_to(ROOT)} "
                f"(copy {path.stem}.example.json and adjust it to your template)"
            )

    layout_path = REQUIRED_LAYOUTS["resume template layout"]
    if not layout_path.exists() or runtime_context is None:
        return

    layout = load_json(layout_path, report)
    if layout is None:
        return

    experience_ids = {e.get("id") for e in runtime_context.get("experience", [])}
    experience_modes = {e.get("id"): e.get("tailoring_mode") for e in runtime_context.get("experience", [])}
    project_ids = {p.get("id") for p in runtime_context.get("projects", [])}
    project_modes = {p.get("id"): p.get("tailoring_mode") for p in runtime_context.get("projects", [])}

    selected_slots = 0
    for block in layout.get("experience_blocks", []) or []:
        owner_id = block.get("owner_id")
        if owner_id not in experience_ids:
            report.fail(
                f"resume_template_layout.json experience_blocks owner_id {owner_id!r} does not "
                "match any experience id in runtime_resume_context.json"
            )
        elif experience_modes.get(owner_id) != block.get("tailoring_mode"):
            report.fail(
                f"resume_template_layout.json experience_blocks owner_id {owner_id!r} has "
                f"tailoring_mode {block.get('tailoring_mode')!r}, but runtime_resume_context.json "
                f"has {experience_modes.get(owner_id)!r}"
            )

    for block in layout.get("project_blocks", []) or []:
        if block.get("tailoring_mode") == "selected":
            selected_slots += 1
            continue
        owner_id = block.get("owner_id")
        if owner_id not in project_ids:
            report.fail(
                f"resume_template_layout.json project_blocks owner_id {owner_id!r} does not "
                "match any project id in runtime_resume_context.json"
            )
        elif project_modes.get(owner_id) != block.get("tailoring_mode"):
            report.fail(
                f"resume_template_layout.json project_blocks owner_id {owner_id!r} has "
                f"tailoring_mode {block.get('tailoring_mode')!r}, but runtime_resume_context.json "
                f"has {project_modes.get(owner_id)!r}"
            )

    if selected_slots != 1:
        report.fail(
            "resume_template_layout.json must define exactly one project_blocks entry "
            f"with tailoring_mode 'selected'; found {selected_slots}"
        )

    # Best-effort structural check of the actual .docx, if python-docx is
    # importable. Never a network/API call.
    template_path = REQUIRED_TEMPLATES["resume template"]
    if template_path.exists():
        try:
            from docx import Document  # type: ignore

            doc = Document(str(template_path))
            min_count = layout.get("min_paragraph_count", 0)
            if len(doc.paragraphs) < min_count:
                report.fail(
                    f"resume_template.docx has {len(doc.paragraphs)} paragraphs; "
                    f"layout requires at least {min_count}"
                )
        except ImportError:
            report.warn("python-docx not installed -- skipped structural .docx paragraph-count check")
        except Exception as exc:
            report.fail(f"Could not open resume_template.docx: {exc}")


def main() -> int:
    report = Report()

    check_files_exist(report)
    runtime_context = check_runtime_context(report)

    if runtime_context is not None:
        check_project_selection_pool(report, runtime_context)
        check_metrics_registry(report, runtime_context)
        check_layouts(report, runtime_context)
    else:
        for label, path in {**REQUIRED_TEMPLATES, **REQUIRED_LAYOUTS}.items():
            if not path.exists():
                report.fail(f"Missing {label}: {path.relative_to(ROOT)}")

    print("=" * 64)
    print("CANDIDATE SETUP VALIDATION")
    print("=" * 64)
    print()

    if report.warnings:
        print(f"WARNINGS ({len(report.warnings)}):")
        for w in report.warnings:
            print(f"  - {w}")
        print()

    if report.failures:
        print(f"FAILURES ({len(report.failures)}):")
        for f in report.failures:
            print(f"  - {f}")
        print()
        print("STATUS: FAIL")
        print()
        print("See INSTRUCTIONS.md and docs/candidate-data-guide.md.")
        return 1

    print("STATUS: PASS")
    if report.warnings:
        print("(warnings above are non-blocking, but worth reviewing)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
