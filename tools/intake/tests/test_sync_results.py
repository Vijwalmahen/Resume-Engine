#!/usr/bin/env python3
"""
Acceptance 12: post-batch reconciliation updates Intake and nothing else.
"""

from __future__ import annotations

import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from support import TempDirTestCase, block_network  # noqa: E402

block_network()

import intake_tracker  # noqa: E402
import sync_results  # noqa: E402

GOOD = "linkedin-4012345678"
FAILED = "greenhouse-examplecorp-7654321"

FINDING_TEXT = "candidate bullet text that must never leave the job directory"


def metrics(batch_id="batch-synthetic-regression"):
    return {
        "batch_id": batch_id,
        "jobs_discovered": 2, "jobs_succeeded": 1, "jobs_failed": 1,
        "jobs": [
            {"job": GOOD, "status": "RESUME_READY",
             "quality_report": {"path": f"output/{GOOD}/resume_quality_report.json",
                                "diagnostic_state": "REVIEW", "artifact_stage": "finalized",
                                "advisory_only": True}},
            {"job": FAILED, "status": "LOCAL_QUALITY_GATE_FAILED",
             "error": "density guard rejected the rendered resume"},
        ],
    }


class SyncTestCase(TempDirTestCase):
    def setUp(self):
        super().setUp()
        self.make_tracker()
        self.metrics_path = self.tmp / "batch_metrics.json"
        self.metrics_path.write_text(json.dumps(metrics()), encoding="utf-8")

        for job_key in (GOOD, FAILED):
            (self.output / job_key).mkdir(parents=True, exist_ok=True)
        (self.output / GOOD / "resume.pdf").write_bytes(b"%PDF-1.7 synthetic")
        (self.output / GOOD / "cover_letter.pdf").write_bytes(b"%PDF-1.7 synthetic")
        (self.output / GOOD / "resume_quality_report.json").write_text(json.dumps({
            "diagnostic_state": "REVIEW",
            "artifact": {"stage": "finalized", "sha256": "0" * 64},
            "findings": [{"severity": "medium", "detail": FINDING_TEXT}],
            "advisory_only": True,
        }), encoding="utf-8")

        intake_tracker.upsert_rows(self.tracker, [
            {"Job Key": GOOD, "Source": "linkedin", "Intake Result": "WRITTEN",
             "Normalized File": f"jobs/batch_inbox/{GOOD}.md",
             "Canonical URL": "https://www.linkedin.com/jobs/view/4012345678",
             "First Seen At": "2026-10-18T00:00:00+00:00"},
            {"Job Key": FAILED, "Source": "greenhouse", "Intake Result": "WRITTEN",
             "Normalized File": f"jobs/batch_inbox/{FAILED}.md",
             "First Seen At": "2026-10-18T00:00:00+00:00"},
        ])

    def sync(self):
        return sync_results.sync(self.metrics_path, self.tracker, self.output)

    def rows(self):
        return {r["Job Key"]: r for r in intake_tracker.read_intake_rows(self.tracker)}


class TestSyncUpdatesIntake(SyncTestCase):
    def test_pipeline_status_batch_id_and_artifacts_are_recorded(self):
        result = self.sync()
        self.assertEqual(result["rows_synced"], 2)
        rows = self.rows()
        self.assertEqual(rows[GOOD]["Pipeline Status"], "RESUME_READY")
        self.assertEqual(rows[GOOD]["Batch ID"], "batch-synthetic-regression")
        self.assertTrue(rows[GOOD]["Final Resume"].endswith(f"{GOOD}/resume.pdf"))
        self.assertTrue(rows[GOOD]["Final Cover Letter"].endswith(f"{GOOD}/cover_letter.pdf"))
        self.assertTrue(rows[GOOD]["Last Sync At"])

    def test_quality_state_and_stage_are_recorded_but_not_the_findings(self):
        self.sync()
        rows = self.rows()
        self.assertEqual(rows[GOOD]["Quality Status"], "REVIEW")
        self.assertEqual(rows[GOOD]["Quality Artifact Stage"], "finalized")
        self.assertIn("resume_quality_report.json", rows[GOOD]["Quality Report"])
        for row in rows.values():
            for value in row.values():
                self.assertNotIn(FINDING_TEXT, value)

    def test_a_failed_job_records_its_status_and_a_short_detail(self):
        self.sync()
        row = self.rows()[FAILED]
        self.assertEqual(row["Pipeline Status"], "LOCAL_QUALITY_GATE_FAILED")
        self.assertEqual(row["Error Code"], "LOCAL_QUALITY_GATE_FAILED")
        self.assertIn("density guard", row["Error Detail"])
        self.assertLessEqual(len(row["Error Detail"]), sync_results.ERROR_DETAIL_MAX_CHARS)
        self.assertEqual(row["Final Resume"], "", "no artifact path for a job that produced none")

    def test_sync_is_idempotent(self):
        self.sync()
        first = self.rows()
        self.sync()
        second = self.rows()
        self.assertEqual(set(first), set(second))
        self.assertEqual(first[GOOD]["Pipeline Status"], second[GOOD]["Pipeline Status"])
        self.assertEqual(len(second), 2)

    def test_a_job_with_no_intake_row_is_reported_not_invented(self):
        payload = metrics()
        payload["jobs"].append({"job": "lever-examplecorp-unknown", "status": "RESUME_READY"})
        self.metrics_path.write_text(json.dumps(payload), encoding="utf-8")
        result = self.sync()
        self.assertEqual(result["unmatched_job_keys"], ["lever-examplecorp-unknown"])
        self.assertNotIn("lever-examplecorp-unknown", self.rows())


class TestSyncTouchesNothingElse(SyncTestCase):
    def applications(self):
        from openpyxl import load_workbook

        wb = load_workbook(self.tracker)
        try:
            return [list(r) for r in wb[intake_tracker.APPLICATIONS_SHEET].iter_rows(values_only=True)]
        finally:
            wb.close()

    def test_application_status_and_notes_are_never_written(self):
        from openpyxl import load_workbook

        wb = load_workbook(self.tracker)
        ws = wb[intake_tracker.APPLICATIONS_SHEET]
        index = {h: i for i, h in enumerate(intake_tracker.APPLICATIONS_HEADERS)}
        row = [""] * len(intake_tracker.APPLICATIONS_HEADERS)
        row[index["Job Key"]] = GOOD
        row[index["Application Status"]] = "APPLIED"
        row[index["Applied Date"]] = "2026-10-12"
        row[index["Notes"]] = "my own note"
        ws.append(row)

        intake_ws = wb[intake_tracker.INTAKE_SHEET]
        notes_col = intake_tracker.INTAKE_HEADERS.index("Notes") + 1
        intake_ws.cell(row=2, column=notes_col, value="my intake note")
        wb.save(self.tracker)
        wb.close()

        before = self.applications()
        self.sync()
        self.assertEqual(self.applications(), before)
        self.assertEqual(self.rows()[GOOD]["Notes"], "my intake note")

    def test_sync_cannot_write_outside_its_own_columns(self):
        self.sync()
        rows = self.rows()
        self.assertEqual(rows[GOOD]["Source"], "linkedin")
        self.assertEqual(rows[GOOD]["Intake Result"], "WRITTEN")
        self.assertEqual(rows[GOOD]["First Seen At"], "2026-10-18T00:00:00+00:00")
        self.assertEqual(
            intake_tracker.SYNC_OWNED_HEADERS & {"Job Key", "Source", "Intake Result", "Notes"},
            set(),
        )

    def test_it_refuses_a_workbook_with_a_foreign_schema(self):
        from openpyxl import Workbook

        foreign = self.tracking / "foreign.xlsx"
        wb = Workbook()
        wb.active.title = "Applications"
        wb.active.append(["Job Key", "Surprise"])
        wb.save(foreign)
        wb.close()
        before = foreign.read_bytes()
        with self.assertRaises(intake_tracker.IntakeSchemaError):
            sync_results.sync(self.metrics_path, foreign, self.output)
        self.assertEqual(foreign.read_bytes(), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
