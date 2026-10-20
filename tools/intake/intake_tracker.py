#!/usr/bin/env python3
"""
The `Intake` worksheet: one row per inspected posting, successful or not.

This writer lives *beside* the V8.6 production tracker, never on top of it.
The workbook it opens is the same one `tools/resume/v86/tracker_writer.py`
owns (`tracking/job_application_tracker.xlsx`), and the rules are:

  - The 23-column `Applications` sheet (config/application_tracker_schema.md)
    is READ for validation and never written. Its headers, their order, and
    every existing row are left byte-for-byte alone.
  - A workbook whose `Applications` sheet does not match that schema — or an
    `Intake` sheet that does not match this module's own — produces a clear
    refusal (`IntakeSchemaError`), never a destructive rewrite. This is the
    same failure mode the dormant 35-column writer in
    `orchestrator/tracker_xlsx.py` now guards against, from the other side.
  - Writes take the SAME lock the production writer takes
    (`tracking/.job_application_tracker.xlsx.lock`, via
    `tools/resume/v86/common.file_lock`), so an intake write and a pipeline
    write can never interleave on one workbook.
  - The save is atomic: openpyxl writes a temp file in the destination
    directory, which is then `os.replace`d onto the real path.
  - `Notes` is the user's. It is written exactly once, as an empty string, on
    the row-creation path, and never touched again by anything in this
    package. `Application Status`, `Applied Date` and `Notes` on the
    `Applications` sheet are not this module's business at all.

Rows are upserted by `Job Key`, which is the same key the posting's
`jobs/batch_inbox/<job_key>.md` file is named after and the same key the
pipeline later writes into `Applications` — see `dedup.py`.
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path
from typing import Optional

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools" / "resume" / "v86"))
from common import file_lock  # noqa: E402
import tracker_writer  # noqa: E402

INTAKE_SHEET = "Intake"
APPLICATIONS_SHEET = "Applications"

# One source of truth: the production writer's own column list.
APPLICATIONS_HEADERS = list(tracker_writer.HEADERS)

# The 32-column Intake schema: source identity, scrape status, description
# length, normalized file path, pipeline status, output paths, quality status,
# error details and timestamps.
INTAKE_HEADERS = [
    "Job Key",                 # 1  identity (== batch_inbox file stem)
    "Source",                  # 2
    "Source Job ID",           # 3
    "Canonical URL",           # 4
    "Application URL",         # 5
    "Company",                 # 6
    "Role",                    # 7
    "Location",                # 8
    "Workplace Type",          # 9
    "Employment Type",         # 10
    "Date Posted",             # 11
    "Compensation",            # 12
    "URL Status",              # 13 SUPPORTED / INCOMPLETE / UNSUPPORTED
    "Scrape Status",           # 14 NOT_ATTEMPTED / SCRAPED / SCRAPE_FAILED / SKIPPED_DUPLICATE
    "Completeness Status",     # 15 COMPLETE / BLOCKED / INCOMPLETE_DESCRIPTION / MALFORMED_RESULT
    "Description Length",      # 16
    "Duplicate Decision",      # 17 NEW / DUPLICATE_*
    "Normalized File",         # 18 jobs/batch_inbox/<job_key>.md
    "Intake Result",           # 19 WRITTEN / REJECTED / DUPLICATE / DRY_RUN
    "Pipeline Status",         # 20 filled by sync_results.py from batch metrics
    "Batch ID",                # 21
    "Final Resume",            # 22
    "Final Cover Letter",      # 23
    "Quality Status",          # 24 the advisory diagnostic_state
    "Quality Artifact Stage",  # 25
    "Quality Report",          # 26 path only; never the findings
    "Error Code",              # 27 why this row is not progressing
    "Error Detail",            # 28 one short line; never report findings
    "First Seen At",           # 29 set once, on row creation
    "Last Intake At",          # 30
    "Last Sync At",            # 31 set by sync_results.py only
    "Notes",                   # 32 THE USER'S. Never written after creation.
]
assert len(INTAKE_HEADERS) == 32

# Written once at row creation and never again by this package.
USER_OWNED_INTAKE_HEADERS = {"Notes"}

# Set on creation only, so a re-intake never rewrites discovery history.
CREATE_ONLY_HEADERS = {"First Seen At"}

# Columns sync_results.py is allowed to write (it must never touch identity,
# intake outcome, or anything the user owns).
SYNC_OWNED_HEADERS = {
    "Pipeline Status", "Batch ID", "Final Resume", "Final Cover Letter",
    "Quality Status", "Quality Artifact Stage", "Quality Report",
    "Error Code", "Error Detail", "Last Sync At",
}


class IntakeSchemaError(Exception):
    """Raised instead of modifying a workbook whose schema is not what this
    module expects."""


def lock_path_for(workbook_path) -> Path:
    """The production writer's lock convention, reproduced exactly:
    `<dir>/.<workbook name>.lock` — see tools/resume/v86/tracker_writer.py."""
    workbook_path = Path(workbook_path)
    return workbook_path.parent / f".{workbook_path.name}.lock"


def _header_row(ws) -> list:
    headers = [("" if c.value is None else str(c.value).strip()) for c in ws[1]]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def _assert_applications_schema(wb, path) -> None:
    if APPLICATIONS_SHEET not in wb.sheetnames:
        raise IntakeSchemaError(
            f"{path} has no {APPLICATIONS_SHEET!r} worksheet (sheets: {wb.sheetnames}). "
            "Intake only extends a workbook that already holds the production "
            "Applications tracker; it never creates one inside an unknown workbook."
        )
    headers = _header_row(wb[APPLICATIONS_SHEET])
    if headers != APPLICATIONS_HEADERS:
        missing = [h for h in APPLICATIONS_HEADERS if h not in headers]
        unexpected = [h for h in headers if h not in APPLICATIONS_HEADERS]
        raise IntakeSchemaError(
            f"{path}'s {APPLICATIONS_SHEET!r} sheet has a {len(headers)}-column header row that "
            f"does not match the {len(APPLICATIONS_HEADERS)}-column production schema "
            f"(config/application_tracker_schema.md). Refusing to modify this workbook. "
            f"missing={missing} unexpected={unexpected}"
        )


def _assert_intake_schema(wb, path) -> None:
    if INTAKE_SHEET not in wb.sheetnames:
        return
    headers = _header_row(wb[INTAKE_SHEET])
    if headers != INTAKE_HEADERS:
        raise IntakeSchemaError(
            f"{path}'s {INTAKE_SHEET!r} sheet has a {len(headers)}-column header row that does not "
            f"match this module's {len(INTAKE_HEADERS)}-column schema. Refusing to modify it."
        )


def _style_header(ws) -> None:
    from openpyxl.styles import Font

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def create_workbook_with_applications(path) -> None:
    """Create a brand-new workbook holding an empty production `Applications`
    sheet, in the production column order. Used only when the target workbook
    does not exist yet."""
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = APPLICATIONS_SHEET
    ws.append(APPLICATIONS_HEADERS)
    _style_header(ws)
    _atomic_save(wb, path)


def _open(path):
    """Open an existing workbook and validate both schemas. Never creates."""
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        raise IntakeSchemaError(
            f"{path} does not exist. Create it first (create_workbook_with_applications) "
            "or point intake at the real tracker."
        )
    wb = load_workbook(path)
    _assert_applications_schema(wb, path)
    _assert_intake_schema(wb, path)
    return wb


def _atomic_save(wb, path) -> None:
    path = Path(path)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}-", suffix=".tmp", dir=str(path.parent))
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def _ensure_intake_sheet(wb):
    if INTAKE_SHEET in wb.sheetnames:
        return wb[INTAKE_SHEET]
    ws = wb.create_sheet(INTAKE_SHEET)
    ws.append(INTAKE_HEADERS)
    _style_header(ws)
    return ws


def assert_compatible(path) -> None:
    """Validate a workbook's `Applications` (and `Intake`, when present)
    schema without modifying anything. Read-only counterpart of the guard
    every write already goes through — for callers that want to fail loudly
    before doing any work."""
    wb = _open(path)
    wb.close()


def read_intake_rows(path) -> list:
    """Every Intake row as {header: value} dicts, in sheet order. Returns []
    when the workbook or the sheet does not exist yet."""
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if INTAKE_SHEET not in wb.sheetnames:
            return []
        ws = wb[INTAKE_SHEET]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [("" if h is None else str(h).strip()) for h in rows[0]]
        out = []
        for values in rows[1:]:
            row = {h: ("" if v is None else str(v)) for h, v in zip(headers, values) if h}
            if (row.get("Job Key") or "").strip():
                out.append(row)
        return out
    finally:
        wb.close()


def upsert_rows(path, rows, *, allowed_headers: Optional[set] = None) -> dict:
    """Create-or-update Intake rows, keyed by `Job Key`, under the shared lock.

    `rows` is a list of {header: value} dicts; only headers present in a dict
    are written, so a partial update (e.g. just the sync-owned columns) leaves
    everything else exactly as it was. Returns a small summary — never the row
    contents.
    """
    path = Path(path)
    created, updated, skipped = [], [], []

    with file_lock(lock_path_for(path)):
        wb = _open(path)
        ws = _ensure_intake_sheet(wb)
        col = {h: i + 1 for i, h in enumerate(INTAKE_HEADERS)}

        row_by_key = {}
        for row in ws.iter_rows(min_row=2):
            key = row[0].value
            if key is not None and str(key).strip():
                row_by_key.setdefault(str(key).strip(), row[0].row)

        for payload in rows:
            job_key = str(payload.get("Job Key") or "").strip()
            if not job_key:
                skipped.append("<missing Job Key>")
                continue

            unknown = [h for h in payload if h not in col]
            if unknown:
                raise IntakeSchemaError(f"unknown Intake column(s): {unknown}")
            if allowed_headers is not None:
                outside = [h for h in payload if h != "Job Key" and h not in allowed_headers]
                if outside:
                    raise IntakeSchemaError(
                        f"this caller may only write {sorted(allowed_headers)}; refused: {outside}"
                    )

            target = row_by_key.get(job_key)
            if target is None:
                ws.append([""] * len(INTAKE_HEADERS))
                target = ws.max_row
                row_by_key[job_key] = target
                ws.cell(row=target, column=col["Job Key"], value=job_key)
                created.append(job_key)
                is_new = True
            else:
                updated.append(job_key)
                is_new = False

            for header, value in payload.items():
                if header == "Job Key":
                    continue
                if header in USER_OWNED_INTAKE_HEADERS:
                    continue  # the user's column; never written here
                if header in CREATE_ONLY_HEADERS and not is_new:
                    continue
                if value is None:
                    continue
                ws.cell(row=target, column=col[header], value=value)

        ws.auto_filter.ref = ws.dimensions
        _atomic_save(wb, path)

    return {"created": created, "updated": updated, "skipped": skipped,
            "workbook": str(path), "sheet": INTAKE_SHEET}


def upsert_row(path, row) -> dict:
    return upsert_rows(path, [row])
