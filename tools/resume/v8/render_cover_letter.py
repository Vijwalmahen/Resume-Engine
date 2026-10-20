#!/usr/bin/env python3
"""
Generic, layout-driven cover-letter renderer.

Mirrors tools/resume/render_resume.py: paragraph indices, which
paragraphs are "static" (identity/signature block, verified unchanged
after editing), and which experience(s) get a dedicated paragraph, all
come from a user-supplied layout file
(reference/cover_letter_template_layout.json; see
reference/cover_letter_template_layout.example.json). Nothing here
hardcodes any candidate's employer or identity.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]

DEFAULT_TEMPLATE = ROOT / "reference" / "cover_letter_template.docx"
DEFAULT_LAYOUT = ROOT / "reference" / "cover_letter_template_layout.json"


def replace_paragraph_text(paragraph, text):
    """Replace text while preserving the formatting of the paragraph's
    first existing run."""

    if paragraph.runs:
        first = paragraph.runs[0]

        font_name = first.font.name
        font_size = first.font.size
        bold = first.bold
        italic = first.italic
        underline = first.underline

        for run in list(paragraph.runs):
            paragraph._element.remove(run._element)

        new_run = paragraph.add_run(text)

        if font_name is not None:
            new_run.font.name = font_name
        if font_size is not None:
            new_run.font.size = font_size
        if bold is not None:
            new_run.bold = bold
        if italic is not None:
            new_run.italic = italic
        if underline is not None:
            new_run.underline = underline

    else:
        paragraph.text = text


def load_layout(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def validate_payload(data: dict, layout: dict) -> None:
    required = ["company", "role", "salutation", "opening", "project_paragraph", "closing"]

    missing = [key for key in required if not str(data.get(key, "")).strip()]

    if missing:
        raise ValueError(f"Missing/empty fields: {missing}")

    for key in ("opening", "project_paragraph", "closing"):
        if "\n" in str(data[key]).strip():
            raise ValueError(f"{key} must be one paragraph")

    experience_paragraphs = data.get("experience_paragraphs")

    if not isinstance(experience_paragraphs, dict):
        raise ValueError("experience_paragraphs must be an object")

    for block in layout.get("experience_paragraph_blocks", []):
        owner_id = block["owner_id"]
        text = experience_paragraphs.get(owner_id)

        if not isinstance(text, str) or not text.strip():
            raise ValueError(
                f"experience_paragraphs[{owner_id}] must be a non-empty string"
            )

        if "\n" in text.strip():
            raise ValueError(f"experience_paragraphs[{owner_id}] must be one paragraph")


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument("--payload", required=True)
    parser.add_argument("--output-docx", required=True)
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    parser.add_argument("--layout", default=str(DEFAULT_LAYOUT))

    args = parser.parse_args()

    payload_path = Path(args.payload)
    output_path = Path(args.output_docx)
    template = Path(args.template)
    layout_path = Path(args.layout)

    data = json.loads(payload_path.read_text())

    if not layout_path.exists():
        raise SystemExit(
            f"Cover-letter template layout not found: {layout_path}\n"
            "Copy reference/cover_letter_template_layout.example.json to "
            f"{layout_path} and adjust it to match your own template."
        )

    layout = load_layout(layout_path)

    validate_payload(data, layout)

    if not template.exists():
        raise SystemExit(f"Template missing: {template}")

    doc = Document(template)
    min_count = layout.get("min_paragraph_count", 0)

    if len(doc.paragraphs) < min_count:
        raise SystemExit(
            "Cover-letter template structure changed: "
            f"expected >={min_count} paragraphs, got {len(doc.paragraphs)}"
        )

    replacements = {
        layout["salutation_index"]: data["salutation"].strip(),
        layout["opening_index"]: data["opening"].strip(),
        layout["project_paragraph_index"]: data["project_paragraph"].strip(),
        layout["closing_index"]: data["closing"].strip(),
    }

    experience_paragraphs = data["experience_paragraphs"]

    for block in layout.get("experience_paragraph_blocks", []):
        replacements[block["index"]] = experience_paragraphs[block["owner_id"]].strip()

    # Capture candidate-specific static template content before replacing
    # dynamic job-specific paragraphs, so we can verify it is unchanged
    # afterward without hardcoding any private identity here.
    static_indices = layout.get("static_indices", [])
    expected_static = {
        index: doc.paragraphs[index].text.strip() for index in static_indices
    }

    for index, value in replacements.items():
        replace_paragraph_text(doc.paragraphs[index], value)

    for index, expected in expected_static.items():
        actual = doc.paragraphs[index].text.strip()

        if actual != expected:
            raise SystemExit(
                f"Template static content changed at paragraph {index}: {actual!r}"
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    print("Cover letter rendered:")
    print(output_path)


if __name__ == "__main__":
    main()
